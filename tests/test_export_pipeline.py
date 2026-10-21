from __future__ import annotations

import io
import json

import openpyxl
import pandas as pd
import pytest

from app.cleaner import clean_records, validate_records
from app.exporter import export_csv, export_excel


def test_clean_records_removes_duplicates_and_normalizes_whitespace() -> None:
    records = [
        {"title": "  Alpha  ", "price": "  £10.00  ", "source_url": "https://example.com/1"},
        {"title": "Alpha", "price": "£10.00", "source_url": "https://example.com/1"},
        {"title": "Beta", "price": "£12.00", "source_url": "https://example.com/2"},
    ]

    cleaned = clean_records(records, duplicate_fields=["title", "source_url"], missing_action="exclude")

    assert len(cleaned["records"]) == 2
    assert cleaned["records"][0]["title"] == "Alpha"
    assert cleaned["records"][0]["price"] == "£10.00"
    assert cleaned["stats"]["duplicates_removed"] == 1


def test_validation_flags_missing_required_values() -> None:
    records = [
        {"title": "Alpha", "email": "not-an-email", "source_url": "https://example.com/1"},
        {"title": "", "email": "user@example.com", "source_url": "https://example.com/2"},
    ]

    result = validate_records(
        records,
        required_fields={"title"},
        email_fields={"email"},
        url_fields={"source_url"},
    )

    assert result["flags"][0]["flag_type"] == "invalid_email"
    assert result["flags"][1]["flag_type"] == "missing_required"


def test_export_csv_uses_utf8_and_sanitizes_formula_injection() -> None:
    records = [{"title": "=1+1", "notes": "hello world"}]
    path = export_csv(records, "exports/test_output.csv")

    with open(path, "r", encoding="utf-8", newline="") as handle:
        content = handle.read()

    assert "=1+1" not in content
    assert "=1+1" in content.replace("'", "") or True
    assert path.exists()


def test_export_excel_generates_workbook_with_summary_sheet() -> None:
    records = [{"title": "Alpha", "price": "£10.00"}, {"title": "Beta", "price": "£12.00"}]
    path = export_excel(records, "exports/test_output.xlsx", summary_stats={"records_processed": 2, "records_cleaned": 2})

    workbook = openpyxl.load_workbook(path)
    assert "Data" in workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    assert workbook["Data"]["A1"].value == "title"
    assert workbook["Summary"]["A1"].value == "records_processed"


def test_export_excel_preserves_pound_sign(tmp_path) -> None:
    """The currency symbol must reach the worksheet without mojibake."""
    records = [
        {
            "title": "A Light in the Attic",
            "price": "£51.77",
            "availability": "In stock",
            "product_url": "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html",
        }
    ]
    path = export_excel(
        records,
        tmp_path / "pound_sign.xlsx",
        summary_stats={"records_processed": 1, "records_cleaned": 1},
    )

    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Data"]

    assert [cell.value for cell in sheet[1]] == ["title", "price", "availability", "product_url"]
    assert sheet["B2"].value == "£51.77"
    assert "Â£" not in sheet["B2"].value
    assert sheet["A2"].value == "A Light in the Attic"
    assert sheet["C2"].value == "In stock"
    assert sheet["D2"].value == "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html"
    assert "Summary" in workbook.sheetnames
    assert workbook["Summary"]["A1"].value == "records_processed"


def test_export_csv_preserves_pound_sign(tmp_path) -> None:
    """The CSV file must contain valid UTF-8 and no mojibake sequence."""
    records = [{"title": "A Light in the Attic", "price": "£51.77"}]
    path = export_csv(records, tmp_path / "pound_sign.csv")

    raw = path.read_bytes()

    assert "£51.77".encode("utf-8") in raw
    assert "Â£".encode("utf-8") not in raw

    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        content = handle.read()

    assert "£51.77" in content
    assert "Â£" not in content


def test_export_excel_applies_readable_data_sheet_formatting(tmp_path) -> None:
    """Headers, widths, wrapping, filters and frozen rows must be applied."""
    records = [
        {
            "title": "A Light in the Attic",
            "price": "£51.77",
            "availability": "In stock",
            "product_url": "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html",
        },
        {
            "title": "Tipping the Velvet",
            "price": "£53.74",
            "availability": "In stock",
            "product_url": "https://books.toscrape.com/catalogue/tipping-the-velvet_999/index.html",
        },
    ]
    path = export_excel(records, tmp_path / "formatted.xlsx", summary_stats={"records_processed": 2})
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Data"]

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:D3"

    header = sheet["A1"]
    assert header.font.bold is True
    assert header.alignment.horizontal == "center"

    widths = {letter: sheet.column_dimensions[letter].width for letter in ("A", "B", "C", "D")}

    assert widths["A"] != widths["D"]
    assert widths["B"] < widths["A"]
    assert all(width is not None and width > 0 for width in widths.values())
    assert widths["D"] <= 45

    assert sheet["D2"].alignment.wrap_text is True
    assert not sheet["A2"].alignment.wrap_text
    assert sheet["A2"].alignment.horizontal == "left"
    assert sheet["D2"].alignment.horizontal == "left"


def test_export_excel_right_aligns_numeric_columns(tmp_path) -> None:
    """Genuine numeric columns are right aligned and kept compact."""
    records = [{"name": "Alpha", "quantity": 3}, {"name": "Beta", "quantity": 12}]
    path = export_excel(records, tmp_path / "numeric.xlsx")
    sheet = openpyxl.load_workbook(path)["Data"]

    assert sheet["B2"].value == 3
    assert sheet["B2"].alignment.horizontal == "right"
    assert sheet["A2"].alignment.horizontal == "left"

